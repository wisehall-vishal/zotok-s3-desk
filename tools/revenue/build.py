import json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from core import key, fuzzy

R = json.load(open('/home/claude/rev/results.json'))
import os
OUT = os.environ.get('REV_OUT', '/home/claude/rev/ZoTok_MoM_Revenue_Jan-Sep26.xlsx')
MONTHS = ['2026-01','2026-02','2026-03','2026-04','2026-05','2026-06','2026-07','2026-08','2026-09']
LBL = ['Jan-26','Feb-26','Mar-26','Apr-26','May-26','Jun-26','Jul-26','Aug-26','Sep-26']
REGM = MONTHS[3:]

F = 'Arial'
H1 = Font(name=F, size=14, bold=True); H2 = Font(name=F, size=11, bold=True)
HD = Font(name=F, size=10, bold=True, color='FFFFFF'); N = Font(name=F, size=10)
BLUE = Font(name=F, size=10, color='0000FF'); GRN = Font(name=F, size=10, color='008000')
NOTE = Font(name=F, size=9, italic=True, color='555555'); BOLD = Font(name=F, size=10, bold=True)
HFILL = PatternFill('solid', fgColor='1F3A3A'); SUB = PatternFill('solid', fgColor='E8F0EE')
FLAG = PatternFill('solid', fgColor='FCE4D6'); OKF = PatternFill('solid', fgColor='E2EFDA')
thin = Side(style='thin', color='CCCCCC'); BOX = Border(bottom=thin)
INR = '#,##0;(#,##0);"–"'; PCT = '0%;(0%);"–"'; LAKH = '#,##0.0,,"L";(#,##0.0,,"L");"–"'

def hdr(ws, row, vals, col=1):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=col + i, value=v); c.font = HD; c.fill = HFILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[L(i)].width = x

wb = Workbook()

# ================= DATA TABS =================
inv = sorted(R['inv_lines'], key=lambda x: (x['date'], x['vch']))
cat_map = {'—': 'Pre-2026 account'}
rep_fix = {'Girish Sir': 'Girish', 'Girish Sir / Vishal': 'Girish/Vishal', 'Bhagysree': 'Bhagyasree', 'Rohit': 'Rohith', 'Ravi': 'Ravi Kishore'}
wi = wb.active; wi.title = 'Invoices'
invname = {}
for x in inv: invname.setdefault(key(x['cust']), x['cust'])
hdr(wi, 1, ['Date', 'Month', 'Customer', 'Voucher', 'Amount incl GST (₹)', 'Revenue ex-GST (₹)', 'Account Manager', 'Closer (S5)', 'Category', 'Billing type', 'I2-I7 stage', 'Customer group'])
for i, x in enumerate(inv, 2):
    from datetime import datetime
    vals = [datetime.strptime(x['date'], '%Y-%m-%d'), x['m'], x['cust'], x['vch'], x['gross'], f'=E{i}/Notes!$B$4', x['am'],
            rep_fix.get(x['rep'], x['rep']), cat_map.get(x['cat'], x['cat']), x['type'], x['stage'], invname[key(x['cust'])]]
    for j, v in enumerate(vals, 1):
        c = wi.cell(row=i, column=j, value=v); c.font = BLUE if j in (1, 3, 4, 5) else N
    wi.cell(row=i, column=1).number_format = 'dd-mmm-yy'
    wi.cell(row=i, column=5).number_format = INR; wi.cell(row=i, column=6).number_format = INR
NI = len(inv) + 1
widths(wi, [11, 9, 40, 16, 15, 15, 16, 14, 16, 26, 30, 36]); wi.freeze_panes = 'A2'; wi.auto_filter.ref = f'A1:L{NI}'

# receipts: map to invoice customer name where possible
wr = wb.create_sheet('Receipts')
hdr(wr, 1, ['Date', 'Month', 'Party (as booked)', 'Amount (₹)', 'New/Old tag', 'Suspense?', 'Matched invoice customer'])
rec = R['rec_lines']
for i, r in enumerate(rec, 2):
    k = key(r['cust']); f = fuzzy(k, invname.keys(), 0.86)
    matched = invname[f] if (f and not r['suspense']) else ('(suspense – unidentified)' if r['suspense'] else '(no invoice Apr–Sep)')
    vals = [datetime.strptime(r['date'], '%Y-%m-%d'), r['m'], r['cust'], r['amt'], r['tag'], 'Y' if r['suspense'] else 'N', matched]
    for j, v in enumerate(vals, 1):
        c = wr.cell(row=i, column=j, value=v); c.font = BLUE if j in (1, 3, 4, 5) else N
    wr.cell(row=i, column=1).number_format = 'dd-mmm-yy'; wr.cell(row=i, column=4).number_format = INR
NR = len(rec) + 1
widths(wr, [11, 9, 42, 14, 10, 10, 40]); wr.freeze_panes = 'A2'; wr.auto_filter.ref = f'A1:G{NR}'

# tracker (I2-I7 monthly schedule) — rebuilt from core
from core import load
_, _, _, i27 = load()
wt = wb.create_sheet('I2-I7 Schedule')
hdr(wt, 1, ['Account', 'Account Manager', 'Closer', 'Stage'] + LBL)
for i, a in enumerate(i27, 2):
    wt.cell(row=i, column=1, value=a['cust']).font = BLUE
    wt.cell(row=i, column=2, value=a['am']).font = BLUE
    wt.cell(row=i, column=3, value=rep_fix.get(a['rep'], a['rep'])).font = BLUE
    wt.cell(row=i, column=4, value=a['stage']).font = BLUE
    for j, m in enumerate(MONTHS):
        c = wt.cell(row=i, column=5 + j, value=a['sched'].get(m, 0) or None); c.font = BLUE; c.number_format = INR
NT = len(i27) + 1
widths(wt, [40, 16, 14, 34] + [11] * 9); wt.freeze_panes = 'E2'; wt.auto_filter.ref = f'A1:M{NT}'

INVR = lambda col: f"Invoices!${col}$2:${col}${NI}"
RECR = lambda col: f"Receipts!${col}$2:${col}${NR}"
TR = lambda col: f"'I2-I7 Schedule'!${col}$2:${col}${NT}"

# ================= SUMMARY =================
ws = wb.create_sheet('MoM Summary', 0)
ws['A1'] = 'ZoTok — Month-on-Month Revenue, Jan–Sep 2026'; ws['A1'].font = H1
ws['A2'] = ('Revenue = Invoice Register ex-GST (accounting source of truth). Collections = Receipt Register (cash, incl GST). '
            'Registers start 1-Apr-26 (FY26-27 voucher series), so Jan–Mar show only Dev\'s I2-I7 billing schedule.')
ws['A2'].font = NOTE; ws.merge_cells('A2:K2'); ws.row_dimensions[2].height = 28; ws['A2'].alignment = Alignment(wrap_text=True)
hdr(ws, 4, ['Metric'] + LBL + ['Apr–Sep'])
for j, m in enumerate(MONTHS):
    ws.cell(row=3, column=2 + j, value=m).font = Font(name=F, size=8, color='999999')
rows = []
def addrow(label, fn, fmt=INR, total='sum', bold=False, note=None, fill=None):
    r = 5 + len(rows); rows.append(label)
    c = ws.cell(row=r, column=1, value=label); c.font = BOLD if bold else N
    if note: c.comment = Comment(note, 'Claude')
    for j, m in enumerate(MONTHS):
        col = L(2 + j); cc = ws.cell(row=r, column=2 + j, value=fn(col, r, j)); cc.number_format = fmt; cc.font = BOLD if bold else N
        if fill: cc.fill = fill
    t = ws.cell(row=r, column=11)
    if total == 'sum': t.value = f'=SUM(E{r}:J{r})'
    elif callable(total): t.value = total(r)
    t.number_format = fmt; t.font = BOLD
    if fill: c.fill = fill; t.fill = fill
    return r

def sec(title):
    r = 5 + len(rows); rows.append(title)
    c = ws.cell(row=r, column=1, value=title); c.font = H2
    for j in range(1, 12): ws.cell(row=r, column=j).fill = SUB
    return r

sec('REVENUE (invoiced)')
rNet = addrow('Revenue ex-GST (₹)', lambda col, r, j: f'=SUMIFS({INVR("F")},{INVR("B")},{col}$3)', bold=True)
rMoM = addrow('  MoM growth', lambda col, r, j: '' if j <= 3 else f'=IFERROR({col}{rNet}/{L(1+j)}{rNet}-1,"")', PCT, total=None)
rNew = addrow('  of which New accounts (setup + 1st bill)', lambda col, r, j: f'=SUMIFS({INVR("F")},{INVR("B")},{col}$3,{INVR("J")},"New*")',
              note='First invoice to a customer within 75 days of its S5 closure date.')
rOld1 = addrow('  of which first invoice this FY, older account', lambda col, r, j: f'=SUMIFS({INVR("F")},{INVR("B")},{col}$3,{INVR("J")},"First*")')
rRep = addrow('  of which Repeat billing (recurring proxy)', lambda col, r, j: f'=SUMIFS({INVR("F")},{INVR("B")},{col}$3,{INVR("J")},"Recurring*")',
              note='2nd+ invoice to the same customer inside Apr–Sep. Closest thing the register has to "MRR actually billed".')
rGross = addrow('Invoiced incl GST (₹)', lambda col, r, j: f'=SUMIFS({INVR("E")},{INVR("B")},{col}$3)')
rNinv = addrow('# invoices', lambda col, r, j: f'=COUNTIFS({INVR("B")},{col}$3)', '#,##0;;"–"')
rNcust = addrow('# customers billed', lambda col, r, j: f'=SUMPRODUCT(({INVR("B")}={col}$3)/COUNTIFS({INVR("B")},{INVR("B")},{INVR("C")},{INVR("C")}))',
                '#,##0;;"–"', total=lambda r: f'=SUMPRODUCT(1/COUNTIF({INVR("C")},{INVR("C")}))')
rAvg = addrow('Avg revenue per customer billed (₹)', lambda col, r, j: f'=IFERROR({col}{rNet}/{col}{rNcust},0)', total=lambda r: f'=IFERROR(K{rNet}/K{rNcust},0)')
rTop1 = addrow('Top customer share of month', lambda col, r, j: '' if j < 3 else f"=IFERROR(MAX('Customer x Month'!{L(2+j-3)}$5:{L(2+j-3)}$400)/{col}{rNet},0)", PCT, total=None)
sec('CASH (collected)')
rCol = addrow('Collected (₹, incl GST)', lambda col, r, j: f'=SUMIFS({RECR("D")},{RECR("B")},{col}$3)', bold=True)
rSus = addrow('  of which Suspense (unidentified payer)', lambda col, r, j: f'=SUMIFS({RECR("D")},{RECR("B")},{col}$3,{RECR("F")},"Y")', fill=FLAG)
rNewC = addrow('  of which tagged New', lambda col, r, j: f'=SUMIFS({RECR("D")},{RECR("B")},{col}$3,{RECR("E")},"New*")')
rCE = addrow('Collection ratio (collected ÷ invoiced incl GST)', lambda col, r, j: f'=IFERROR({col}{rCol}/{col}{rGross},0)', PCT,
             total=lambda r: f'=IFERROR(K{rCol}/K{rGross},0)')
rAR = addrow('Cumulative billed-not-collected since 1-Apr (₹)', lambda col, r, j: '' if j < 3 else f'=SUM($E{rGross}:{col}{rGross})-SUM($E{rCol}:{col}{rCol})',
             total=None, note='No opening balances: Apr receipts against FY25-26 invoices reduce this. Treat as a trend line, not a ledger balance.')
sec('PLAN vs ACTUAL (Dev\'s I2-I7 schedule)')
rTr = addrow('I2-I7 scheduled billing (₹)', lambda col, r, j: f'=SUM({TR(L(5+j))})', bold=True)
rTrN = addrow('  # accounts scheduled', lambda col, r, j: f'=COUNTIF({TR(L(5+j))},">0")', '#,##0;;"–"', total=None)
rGap = addrow('Invoiced ex-GST minus I2-I7 schedule (₹)', lambda col, r, j: '' if j < 3 else f'={col}{rNet}-{col}{rTr}', fill=FLAG,
              total=lambda r: f'=K{rNet}-K{rTr}')
for r in range(5, 5 + len(rows)):
    for j in range(1, 12): ws.cell(row=r, column=j).border = BOX
nr = 5 + len(rows) + 1
ws.cell(row=nr, column=1, value='Grey row 3 holds the month keys the formulas look up. Blue text on data tabs = pasted source data; everything here is a formula.').font = NOTE
widths(ws, [46] + [12] * 10); ws.freeze_panes = 'B5'

# ================= SLICES =================
def slice_tab(name, title, src_col_inv, src_col_tr, members, note):
    s = wb.create_sheet(name)
    s['A1'] = title; s['A1'].font = H1; s['A2'] = note; s['A2'].font = NOTE
    s['A4'] = 'Revenue ex-GST (Invoice Register)'; s['A4'].font = H2
    hdr(s, 5, [name.split(' ')[1] if ' ' in name else 'Group'] + LBL[3:] + ['Apr–Sep', 'Share'])
    for j, m in enumerate(REGM): s.cell(row=4, column=2 + j, value=m).font = Font(name=F, size=8, color='999999')
    r0 = 6
    for i, mem in enumerate(members):
        r = r0 + i; s.cell(row=r, column=1, value=mem).font = N
        for j in range(6):
            c = s.cell(row=r, column=2 + j, value=f'=SUMIFS({INVR("F")},{INVR("B")},{L(2+j)}$4,{INVR(src_col_inv)},$A{r})'); c.number_format = INR; c.font = N
        s.cell(row=r, column=8, value=f'=SUM(B{r}:G{r})').number_format = INR
        s.cell(row=r, column=9, value=f'=IFERROR(H{r}/H${r0+len(members)},0)').number_format = PCT
    rt = r0 + len(members); s.cell(row=rt, column=1, value='Total').font = BOLD
    for j in range(2, 9):
        c = s.cell(row=rt, column=j, value=f'=SUM({L(j)}{r0}:{L(j)}{rt-1})'); c.number_format = INR; c.font = BOLD
    s.cell(row=rt + 1, column=1, value='Check vs MoM Summary (should be 0)').font = NOTE
    c = s.cell(row=rt + 1, column=8, value=f"=H{rt}-'MoM Summary'!K{rNet}"); c.number_format = INR; c.font = NOTE
    if src_col_tr:
        r2 = rt + 3; s.cell(row=r2, column=1, value='I2-I7 scheduled billing, Jan–Sep (plan view)').font = H2
        hdr(s, r2 + 1, ['Group'] + LBL + ['Jan–Sep'])
        for i, mem in enumerate(members):
            r = r2 + 2 + i; s.cell(row=r, column=1, value=mem).font = N
            for j in range(9):
                c = s.cell(row=r, column=2 + j, value=f"=SUMIFS({TR(L(5+j))},{TR(src_col_tr)},$A{r})"); c.number_format = INR; c.font = N
            s.cell(row=r, column=11, value=f'=SUM(B{r}:J{r})').number_format = INR
    widths(s, [28] + [12] * 10); s.freeze_panes = 'B6'
    return s

amt_by = lambda fld: sorted({x[fld] for x in inv}, key=lambda m: -sum(y['net'] for y in inv if y[fld] == m))
ams = sorted({rep_fix.get(x['am'], x['am']) for x in inv} | {a['am'] for a in i27 if any(a['sched'].get(m) for m in MONTHS)},
             key=lambda m: -sum(y['net'] for y in inv if y['am'] == m))
slice_tab('By AccountManager', 'Revenue by Account Manager', 'G', 'B', ams,
          'AM from Dev\'s I2-I7 sheet. "(not in I2-I7)" = invoiced customers with no I2-I7 account — onboarding not logged.')
reps = sorted({rep_fix.get(x['rep'], x['rep']) for x in inv}, key=lambda m: -sum(y['net'] for y in inv if rep_fix.get(y['rep'], y['rep']) == m))
slice_tab('By Closer', 'Revenue by closing salesperson', 'H', 'C', reps, 'Closer from S5 (or I2-I7 where S5 has no 2026 row).')
cats = ['AC2', 'ICP', 'AC1', 'Pre-2026 account']
slice_tab('By Category', 'Revenue by account category', 'I', None, cats, 'Category from S5 contract type. "Pre-2026 account" = closed before Jan-26, so no 2026 S5 row.')
ty = ['New (setup + 1st billing)', 'First invoice in FY (old a/c)', 'Recurring / repeat']
slice_tab('By BillingType', 'Revenue by billing type', 'J', None, ty, 'New = first invoice within 75 days of S5 closure. Recurring = 2nd+ invoice to same customer Apr–Sep.')

# Customer x Month
cm = wb.create_sheet('Customer x Month')
cm['A1'] = 'Revenue ex-GST by customer, Apr–Sep 2026'; cm['A1'].font = H1
cm['A2'] = 'Sorted by Apr–Sep revenue. "Months billed" shows how often a customer is actually invoiced — monthly MRR accounts should read 6.'; cm['A2'].font = NOTE
hdr(cm, 4, ['Customer'] + LBL[3:] + ['Apr–Sep', 'Share', 'Months billed', 'Cumulative share'])
custs = sorted({x['cust'] for x in inv}, key=lambda c: -sum(x['net'] for x in inv if x['cust'] == c))
for j, m in enumerate(REGM): cm.cell(row=3, column=2 + j, value=m).font = Font(name=F, size=8, color='999999')
for i, c_ in enumerate(custs, 5):
    cm.cell(row=i, column=1, value=c_).font = N
    for j in range(6):
        c = cm.cell(row=i, column=2 + j, value=f'=SUMIFS({INVR("F")},{INVR("B")},{L(2+j)}$3,{INVR("C")},$A{i})'); c.number_format = INR; c.font = N
    cm.cell(row=i, column=8, value=f'=SUM(B{i}:G{i})').number_format = INR
    cm.cell(row=i, column=9, value=f"=IFERROR(H{i}/'MoM Summary'!$K${rNet},0)").number_format = PCT
    cm.cell(row=i, column=10, value=f'=COUNTIF(B{i}:G{i},">0")').number_format = '0'
    cm.cell(row=i, column=11, value=f'=SUM($I$5:I{i})').number_format = PCT
widths(cm, [44] + [12] * 6 + [13, 8, 10, 11]); cm.freeze_panes = 'B5'; cm.auto_filter.ref = f'A4:K{4+len(custs)}'

# Receivables
ar = wb.create_sheet('Receivables')
ar['A1'] = 'Billed vs collected by customer, 1-Apr to 30-Sep-26 (incl GST)'; ar['A1'].font = H1
ar['A2'] = 'Grouped by customer (all legal-entity spellings combined). No opening balances in the registers: a negative balance usually means cash received against a FY25-26 invoice. Sorted by balance.'; ar['A2'].font = NOTE
hdr(ar, 4, ['Customer', 'Account Manager', 'Invoiced', 'Collected', 'Balance', 'Last invoice', 'Days since last invoice', 'Flag'])
led = [l for l in R['ledger'] if l['inv'] > 0]
for l in led: l['name'] = invname.get(l['k'], l['name'])
for i, l in enumerate(led, 5):
    ar.cell(row=i, column=1, value=l['name']).font = N
    ar.cell(row=i, column=2, value=l['am']).font = N
    ar.cell(row=i, column=3, value=f'=SUMIFS({INVR("E")},{INVR("L")},$A{i})').number_format = INR
    ar.cell(row=i, column=4, value=f'=SUMIFS({RECR("D")},{RECR("G")},$A{i})').number_format = INR
    ar.cell(row=i, column=5, value=f'=C{i}-D{i}').number_format = INR
    ar.cell(row=i, column=6, value=f'=MAXIFS({INVR("A")},{INVR("L")},$A{i})'.replace('MAXIFS', '_xlfn.MAXIFS')).number_format = 'dd-mmm-yy'
    ar.cell(row=i, column=7, value=f'=IF(E{i}>1000,Notes!$B$5-F{i},"")').number_format = '0'
    ar.cell(row=i, column=8, value=f'=IF(E{i}>500000,"Big open",IF(AND(E{i}>1000,G{i}>45),"Overdue >45d",IF(E{i}<-1000,"Credit / prior-FY","")))')
widths(ar, [44, 16, 13, 13, 13, 12, 12, 16]); ar.freeze_panes = 'A5'; ar.auto_filter.ref = f'A4:H{4+len(led)}'

# ================= TRIANGULATION =================
t1 = wb.create_sheet('S5 vs Receipts')
t1['A1'] = 'S5 closures (Apr–Sep) — advance per S5 vs cash in Receipt Register'; t1['A1'].font = H1
t1['A2'] = 'Window: receipts 45 days before to 60 days after closure. "Likely in Suspense" = an unidentified receipt of the exact advance amount.'; t1['A2'].font = NOTE
hdr(t1, 4, ['Customer', 'Closer', 'Category', 'Closed', 'S5 advance', 'Received', 'Difference', 'Receipt dates', 'Invoiced (incl GST)', 'Status', 'Suspense hint', 'In I2-I7?'])
order = {'Not in receipts': 0, 'Likely in Suspense': 1, 'Amount differs': 2, 'No advance in S5, none received': 3, 'Matched': 4}
tri = sorted(R['tri_s5'], key=lambda t: (order.get(t['status'], 9), -t['s5_adv']))
for i, t in enumerate(tri, 5):
    vals = [t['cust'], rep_fix.get(t['rep'], t['rep']), t['cat'], t['closed'], t['s5_adv'], t['received'], f'=F{i}-E{i}', t['rec_dates'], t['invoiced'], t['status'], t['suspense_hint'], 'Y' if t['in_i27'] else 'N']
    for j, v in enumerate(vals, 1):
        c = t1.cell(row=i, column=j, value=v); c.font = N
        if j in (5, 6, 7, 9): c.number_format = INR
    if t['status'] != 'Matched':
        for j in range(1, 13): t1.cell(row=i, column=j).fill = FLAG if t['status'] in ('Not in receipts', 'Likely in Suspense') else PatternFill()
widths(t1, [34, 13, 9, 9, 12, 12, 12, 20, 13, 26, 18, 8]); t1.freeze_panes = 'A5'; t1.auto_filter.ref = f'A4:L{4+len(tri)}'

t2 = wb.create_sheet('I2-I7 vs Invoices')
t2['A1'] = 'Dev\'s I2-I7 billing schedule vs actual invoices, Apr–Sep (ex-GST)'; t2['A1'].font = H1
t2['A2'] = 'Sub-accounts of one customer (e.g. 3 Hindware lines) are combined. Sorted by absolute gap. Timing gaps net out over the window; structural gaps don\'t.'; t2['A2'].font = NOTE
hdr(t2, 4, ['I2-I7 account(s)', 'Account Manager', 'Stage', 'Scheduled Apr–Sep', 'Invoiced Apr–Sep', 'Gap (inv − sched)', 'Months scheduled', 'Months invoiced', 'Read'])
agg = defaultdict(lambda: dict(am='', st='', tr=0, inv=0, mt=0, mi=0))
for t in R['tri_i']:
    a = agg[t['cust']]; a['am'] = t['am']; a['st'] = t['stage']; a['tr'] += t['tracker']; a['inv'] += t['invoiced']
    a['mt'] += 1 if t['tracker'] else 0; a['mi'] += 1 if t['invoiced'] else 0
for i, (c_, a) in enumerate(sorted(agg.items(), key=lambda kv: -abs(kv[1]['inv'] - kv[1]['tr'])), 5):
    vals = [c_, a['am'], a['st'], a['tr'], a['inv'], f'=E{i}-D{i}', a['mt'], a['mi'],
            f'=IF(AND(D{i}>0,E{i}=0),"Scheduled, never invoiced",IF(AND(E{i}>0,D{i}=0),"Invoiced, not in schedule",IF(ABS(F{i})<=0.1*MAX(D{i},E{i}),"Agrees (timing only)",IF(F{i}>0,"Invoiced more than plan","Plan ahead of invoicing"))))']
    for j, v in enumerate(vals, 1):
        c = t2.cell(row=i, column=j, value=v); c.font = N
        if j in (4, 5, 6): c.number_format = INR
widths(t2, [44, 18, 10, 14, 14, 14, 10, 10, 26]); t2.freeze_panes = 'A5'; t2.auto_filter.ref = f'A4:I{4+len(agg)}'

# ================= EXCEPTIONS =================
ex = wb.create_sheet('Exceptions')
ex['A1'] = 'Exceptions for Accounts / Dev / AMs to clear'; ex['A1'].font = H1
r = 3
def block(title, header, rowsv, fmts=None):
    global r
    ex.cell(row=r, column=1, value=title).font = H2; r += 1
    hdr(ex, r, header); r += 1
    for rv in rowsv:
        for j, v in enumerate(rv, 1):
            c = ex.cell(row=r, column=j, value=v); c.font = N
            if fmts and j in fmts: c.number_format = fmts[j]
        r += 1
    r += 1
newtag = {x['cust']: x['tag'] for x in rec}
block('1. Cash received with no invoice in the Apr–Sep register (GST exposure if "New")',
      ['Party', 'Received (₹)', 'Tag', 'Read'],
      [[k, v, newtag.get(k, ''), 'Advance taken, tax invoice not raised' if newtag.get(k, '').strip().lower().startswith('new') or not newtag.get(k, '').strip() else 'Likely against FY25-26 invoice']
       for k, v in sorted(R['rec_no_inv'].items(), key=lambda t: -t[1])], {2: INR})
block('2. Suspense receipts (payer not identified) — with likely owner where the amount matches a closure',
      ['Date', 'Amount (₹)', 'Likely payer'],
      [[s['date'], s['amt'], next((t['cust'] for t in R['tri_s5'] if t['suspense_hint'] and s['date'] in t['suspense_hint'] and abs(t['s5_adv'] - s['amt']) < 1), '₹5,760.76 = recurring small-ticket net of TDS' if abs(s['amt'] - 5760.76) < 0.01 else ('Test / rounding entry' if s['amt'] < 5 else ''))]
       for s in R['suspense']], {2: '#,##0.00'})
block('3. Invoiced customers with no I2-I7 account (onboarding not logged by Dev)', ['Customer', 'Invoiced ex-GST Apr–Sep'],
      [[c_, sum(x['net'] for x in inv if x['cust'] == c_)] for c_ in R['inv_not_in_i27']], {2: INR})
block('4. Register hygiene', ['Item', 'Detail'], [
    ['Missing invoice numbers', 'Z/2627/000150, 151, 152 — gap between 31-Aug and 8-Sep. Cancelled, or missing from the export?'],
    ['Cut-off: invoices', 'BirlaNu ₹6,60,800 dated 2-Jun counted in May subtotal; Bhagyalaxmi ₹4,30,700 dated 17-Jul counted in June subtotal. Grand totals tie.'],
    ['Cut-off: receipts', 'BirlaNu ₹6,04,800 and Hindware ₹3,48,000 dated 2-Jul counted in June subtotal. Grand totals tie.'],
    ['Out-of-sequence numbers', 'Invoices 42/43 and receipts 75–78, 108–110, 127/128, 135/138 booked out of date order.'],
    ['Non-18% amounts', 'Some invoices (₹5,000, ₹20,000, ₹4,956, ₹87,792, ₹16,225) are not clean ×1.18 — check if any are non-GST / reverse-charge. Workbook applies 18% uniformly (Notes!B4).'],
])
widths(ex, [46, 16, 40, 40])

# ================= INSIGHTS & NOTES =================
ins = wb.create_sheet('Insights', 1)
ins['A1'] = 'What the accounting registers say — Apr to Sep 2026'; ins['A1'].font = H1
lines = [
 ('Headline', "₹1.30 Cr revenue ex-GST invoiced Apr–Sep (₹1.53 Cr incl GST); ₹1.02 Cr collected (66%). Sep alone is ₹47.8L — 37% of the half-year — driven by Haldiram, JK Cement, Savex and Orient."),
 ('Real recurring run-rate', "Repeat billing (2nd+ invoice to the same customer) is ₹43L over 6 months ≈ ₹7L/month ex-GST. That, not the I2-I7 MRR column, is the recurring revenue the books can see today — vs the ₹1 Cr/month goal."),
 ('MRR is not invoiced monthly', "87 of 119 customers billed Apr–Sep were invoiced in only ONE month; only 7 were billed in 3+ months. Either billing is quarterly/annual-in-advance, or monthly invoices aren't being raised. Ask Accounts + CS which — it decides whether the Dec MRR number is real."),
 ('Concentration', "Top 3 customers (Haldiram, JK Cement, UltraTech) = 30% of Apr–Sep revenue; top 10 = 68%. AC2 accounts = 67% of revenue from a handful of logos. Small-ticket invoices (≤₹35.4k incl GST) are 114 of 175 invoices but only 13% of value."),
 ('Cash risk is three names', "Open balances at 30-Sep: Orient ₹16.9L, Haldiram ₹16.6L, JK Cement ₹14.3L (all incl GST) = ₹47.7L, i.e. almost the entire Sep gap between billing and cash."),
 ('Plan vs books', "Dev's I2-I7 schedule totals ₹1.14 Cr for Apr–Sep vs ₹1.28 Cr actually invoiced for the same accounts, but only 44 of 93 accounts agree within 10%. I2-I7 is a plan, not a billing record: it shows Orient at ₹1.4L vs ₹17.6L invoiced, and Hindware ₹7.4L vs ₹4.0L."),
 ('Orient puzzle — partly settled', "Books show Orient invoiced ₹9L (31-Jul, = the S5 ₹9L setup) + ₹3.6L (19-Aug) + ₹3.6L + ₹1.4L (Sep). So the ₹9L setup is real and billed; I2-I7's ₹1.4L is just the MRR line starting Sep."),
 ('S5 ↔ cash', "Of 82 S5 closures Apr–Sep: 53 advances match a receipt, 11 differ in amount, 8 have no receipt (₹3.6L — but ₹1.96L is Nexibles, paid 1-Oct after the register cut-off, and CoEvolve's ₹14,750 most likely arrived as 'Beenlightened Media Productions' on 18-Sep), and 3 look parked in Suspense (Vedant ₹29.5k, Om Sai Maalikh ₹23.6k, Vivalicious ₹17.7k). Genuinely missing: Tenali Double Horse, Merceley's, OPOS Water (Abhigna), Kaka Tobacco, Fresh World, Rupmaya (Venkat) — ₹1.5L."),
 ('Money in, no invoice', "₹2.4L received from PM Cona, Paras, Powder Pack, Yantrayug and Frumar with no tax invoice in the register — raise invoices to keep GST clean."),
 ('Jan–Mar gap', "Registers start 1-Apr (FY26-27). For a true 9-month view, ask Accounts for the FY25-26 invoice + receipt registers for Jan–Mar; until then Jan–Mar rely on I2-I7 (₹6.2L, ₹5.1L, ₹14.4L)."),
]
for i, (h, t) in enumerate(lines, 3):
    ins.cell(row=i, column=1, value=h).font = BOLD
    c = ins.cell(row=i, column=2, value=t); c.font = N; c.alignment = Alignment(wrap_text=True, vertical='top')
    ins.cell(row=i, column=1).alignment = Alignment(vertical='top'); ins.row_dimensions[i].height = 44
widths(ins, [26, 120])

nt = wb.create_sheet('Notes')
nt['A1'] = 'Sources, assumptions and how to refresh'; nt['A1'].font = H1
nt['A4'] = 'GST divisor (revenue = invoice ÷ this)'; nt['B4'] = 1.18; nt['B4'].font = BLUE; nt['B4'].fill = PatternFill('solid', fgColor='FFFF00')
nt['A5'] = 'As-of date for ageing'; nt['B5'] = datetime(2026, 9, 30); nt['B5'].number_format = 'dd-mmm-yy'; nt['B5'].font = BLUE; nt['B5'].fill = PatternFill('solid', fgColor='FFFF00')
src = ['Raw Invoice Register.csv & Raw Receipt Register.csv — SharePoint 00_Zono Relevant/MIS/Dec26 Bridge/CSV (Accounts team), read 3-Oct-26. Grand totals tie to the register\'s own subtotals.',
       'RAW I2-I7 MRR Data.csv (Dev) — monthly billing schedule columns Jan–Sep-26; X treated as 0.',
       'Raw S5 Closure Data.csv (Anshul) — 2026 closures for advance, category and closer.',
       'Customer names are matched across sources with a normalised key + a legal-entity alias list (e.g. Recykal = Rapidue Technologies, Polaad = Bhagyalaxmi Rolling Mill, VRF = Vikram Roller Flour Mills, India Gate = KRBL) — aliases were confirmed by exact matching advance amounts.',
       'To refresh: paste new register rows into Invoices / Receipts (blue columns), copy the formula cells down, extend month keys if needed. All summary and slice tabs recalculate.']
for i, s in enumerate(src, 7):
    c = nt.cell(row=i, column=1, value=s); c.font = N; c.alignment = Alignment(wrap_text=True); nt.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6); nt.row_dimensions[i].height = 30
widths(nt, [40, 14, 14, 14, 14, 30])

order = ['MoM Summary', 'Insights', 'By AccountManager', 'By Closer', 'By Category', 'By BillingType', 'Customer x Month', 'Receivables',
         'S5 vs Receipts', 'I2-I7 vs Invoices', 'Exceptions', 'Notes', 'Invoices', 'Receipts', 'I2-I7 Schedule']
wb._sheets = [wb[n] for n in order]
wb.active = 0
for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save(OUT)
print('saved', OUT, 'rows: inv', NI - 1, 'rec', NR - 1, 'i27', NT - 1)
